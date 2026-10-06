"""Step 5: build the 7-slide project summary deck from figures/ into docs/status_deck.pptx.

Run from the repo root after scripts/04_figures.py.
"""
from PIL import Image
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN
from pptx.util import Inches, Pt

W, H = Inches(13.333), Inches(7.5)
INK = RGBColor(0x1A, 0x1A, 0x1A)
MUTED = RGBColor(0x5F, 0x5F, 0x5F)
GREEN = RGBColor(0x1B, 0x78, 0x37)
RUST = RGBColor(0xB3, 0x54, 0x1E)
RULE = RGBColor(0xD8, 0xD8, 0xD8)
PANEL = RGBColor(0xF4, 0xF4, 0xF2)

prs = Presentation()
prs.slide_width, prs.slide_height = W, H
BLANK = prs.slide_layouts[6]

M = Inches(0.62)          # left margin
CW = W - 2 * M            # content width


def slide():
    s = prs.slides.add_slide(BLANK)
    for ph in list(s.placeholders):
        ph._element.getparent().remove(ph._element)
    return s


def tb(s, x, y, w, h, align=PP_ALIGN.LEFT):
    box = s.shapes.add_textbox(x, y, w, h)
    tf = box.text_frame
    tf.word_wrap = True
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    tf.paragraphs[0].alignment = align
    return tf


def para(tf, text, size=14, bold=False, color=INK, space_after=6, first=False,
         align=None, italic=False):
    p = tf.paragraphs[0] if first else tf.add_paragraph()
    p.space_after = Pt(space_after)
    if align:
        p.alignment = align
    r = p.add_run()
    r.text = text
    r.font.size = Pt(size)
    r.font.bold = bold
    r.font.italic = italic
    r.font.color.rgb = color
    r.font.name = "Calibri"
    return p


def rich(tf, chunks, size=14, space_after=6, first=False, bullet_color=None):
    """chunks: list of (text, bold, color, italic)."""
    p = tf.paragraphs[0] if first else tf.add_paragraph()
    p.space_after = Pt(space_after)
    for text, bold, color, italic in chunks:
        r = p.add_run()
        r.text = text
        r.font.size = Pt(size)
        r.font.bold = bold
        r.font.italic = italic
        r.font.color.rgb = color
        r.font.name = "Calibri"
    return p


def header(s, kicker, title):
    tf = tb(s, M, Inches(0.42), CW, Inches(0.3))
    para(tf, kicker.upper(), size=11, bold=True, color=GREEN, space_after=3, first=True)
    tf2 = tb(s, M, Inches(0.74), CW, Inches(0.62))
    para(tf2, title, size=27, bold=True, space_after=0, first=True)
    ln = s.shapes.add_shape(1, M, Inches(1.40), CW, Pt(1.1))
    ln.fill.solid()
    ln.fill.fore_color.rgb = RULE
    ln.line.fill.background()
    ln.shadow.inherit = False
    return s


def pic_fit(s, path, x, y, max_w, max_h):
    iw, ih = Image.open(path).size
    scale = min(max_w / iw, max_h / ih)
    w, h = int(iw * scale), int(ih * scale)
    return s.shapes.add_picture(path, int(x + (max_w - w) / 2), int(y + (max_h - h) / 2),
                                width=w, height=h)


def card(s, x, y, w, h, fill=PANEL):
    r = s.shapes.add_shape(1, x, y, w, h)
    r.fill.solid()
    r.fill.fore_color.rgb = fill
    r.line.color.rgb = RULE
    r.line.width = Pt(0.75)
    r.shadow.inherit = False
    return r


def notes(s, text):
    s.notes_slide.notes_text_frame.text = text


# ---------------------------------------------------------------- 1. title
s = slide()
band = s.shapes.add_shape(1, 0, 0, Inches(0.16), H)
band.fill.solid()
band.fill.fore_color.rgb = GREEN
band.line.fill.background()
band.shadow.inherit = False

tf = tb(s, Inches(0.95), Inches(1.92), Inches(11.4), Inches(2.85))
para(tf, "Is anterior/posterior fate primed in chromatin", size=36, bold=True,
     space_after=2, first=True)
para(tf, "before it appears in transcription?", size=36, bold=True, space_after=16)
para(tf, "Feasibility check and RNA baseline on a mouse secondary-palate multiome",
     size=16, color=MUTED, space_after=4)
rich(tf, [("GSE218576  ", False, MUTED, False),
          ("\u00b7  9 libraries  \u00b7  E12.5\u2013E14.5  \u00b7  38,071 cells passing count QC",
           False, MUTED, False)], size=13, space_after=0)

tf = tb(s, Inches(0.95), Inches(6.42), Inches(11.4), Inches(0.4))
para(tf, "Reanalysis of public data from Yan et al. 2024 (Nat Commun 15:821). "
         "No new experiments; all claims below are from their deposited matrices.",
     size=11, color=MUTED, space_after=0, first=True)
notes(s, "Two days of work: verify the data is usable, then build the RNA baseline "
         "that the chromatin question has to beat.")

# ---------------------------------------------------------------- 2. question
s = slide()
header(s, "The question", "Fate must be written somewhere before it is expressed")

tf = tb(s, M, Inches(1.75), Inches(6.3), Inches(4.5))
para(tf, "The secondary palate patterns along its anterior\u2013posterior axis. "
         "Anterior mesenchyme expresses Shox2 and Msx1; posterior expresses Meox2 and Tbx22. "
         "By E14.5 the two territories are transcriptionally distinct.",
     size=15, space_after=14, first=True)
para(tf, "At E12.5 they are not. The cells are there, the fates are not yet declared in RNA.",
     size=15, space_after=14)
rich(tf, [("So: does chromatin know first? ", True, INK, False),
          ("If accessibility at fate-linked loci is already biased in progenitors whose "
           "transcriptome is uninformative, positional fate is primed epigenetically. "
           "If not, the two modalities commit together.", False, INK, False)], size=15,
     space_after=14)
para(tf, "This is answerable because the assay measures both modalities in the same nucleus.",
     size=15, color=MUTED, space_after=0)

c = card(s, Inches(7.35), Inches(1.75), Inches(5.35), Inches(4.5))
tf = tb(s, Inches(7.72), Inches(2.02), Inches(4.6), Inches(4.0))
para(tf, "WHY THIS DATASET", size=11, bold=True, color=GREEN, space_after=10, first=True)
for k, v in [("Same nucleus", "RNA + ATAC jointly, so no cross-modality cell matching is needed."),
             ("Right window", "E12.5 precedes A/P transcriptional identity; E14.5 is fully patterned."),
             ("Real replicates", "2\u20133 libraries per stage, so batch effects are separable from biology."),
             ("Already public", "Processed matrices and fragments are deposited; nothing to re-align.")]:
    rich(tf, [(k + "  ", True, INK, False), (v, False, MUTED, False)], size=13, space_after=11)
notes(s, "The question is developmental; the method is multimodal. That pairing is the point.")

# ---------------------------------------------------------------- 3. day 1
s = slide()
header(s, "Step 1 \u00b7 feasibility", "The stated risk cleared; a different one appeared")
pic_fit(s, "figures/fig1_day1_feasibility.png", M, Inches(1.62), Inches(7.55), Inches(5.35))

x = Inches(8.45)
tf = tb(s, x, Inches(1.68), Inches(4.3), Inches(0.9))
rich(tf, [("Processed files exist. ", True, GREEN, False),
          ("All 9 libraries carry both a filtered matrix and an ATAC fragment file "
           "(cellranger-arc 2.0.0). No re-alignment needed.", False, INK, False)],
     size=13, space_after=0, first=True)

tf = tb(s, x, Inches(2.88), Inches(4.3), Inches(1.30))
rich(tf, [("QC reproduces. ", True, GREEN, False),
          ("Applying the authors' own thresholds keeps 38,071 of 39,669 barcodes (96%). "
           "They report 36,154 after two further gates that need the fragment files \u2014 "
           "the gap is the right size.", False, INK, False)], size=13, space_after=0, first=True)

tf = tb(s, x, Inches(4.42), Inches(4.3), Inches(1.2))
rich(tf, [("But the peak sets do not match. ", True, RUST, False),
          ("Each library was peak-called separately: 67k\u2013140k peaks each, and only "
           "0.08% of peak IDs recur in any other library. The matrices cannot be "
           "concatenated as they stand.", False, INK, False)], size=13, space_after=0, first=True)

c = card(s, x, Inches(5.82), Inches(4.3), Inches(1.20), RGBColor(0xFC, 0xF3, 0xEC))
tf = tb(s, Inches(8.70), Inches(5.99), Inches(3.9), Inches(0.92))
rich(tf, [("Consequence  ", True, RUST, False),
          ("re-quantify all 38k cells against a consensus set of 196,662 merged intervals, "
           "from 12.3 GB of fragments. That is a compute step, not a checkbox.",
           False, INK, False)], size=12, space_after=0, first=True)
notes(s, "Panel c is the finding. Zero of the nine libraries share a peak vocabulary, "
         "so the ATAC half of the project starts with a re-quantification job.")

# ---------------------------------------------------------------- 4. umap
s = slide()
header(s, "Step 2 \u00b7 RNA baseline", "38,071 cells, annotated from marker expression")
pic_fit(s, "figures/fig2_rna_umap.png", M, Inches(1.58), Inches(7.9), Inches(5.5))

x = Inches(8.85)
tf = tb(s, x, Inches(1.70), Inches(3.9), Inches(5.0))
para(tf, "WHAT THE EMBEDDING SHOWS", size=11, bold=True, color=GREEN, space_after=11, first=True)
for k, v in [("The A/P axis is real",
              "Shox2 averages 1.65 in anterior cells vs 0.17 in posterior; Meox2 1.28 vs 0.11."),
             ("Stage structures the map",
              "E12.5 occupies its own territory (panel b) and only separates into A and P from E13.5."),
             ("Osteogenesis follows",
              "Zero osteogenic cells at E12.5, peaking at E14.0 \u2014 consistent with Sp7 onset.")]:
    rich(tf, [(k, True, INK, False)], size=13, space_after=3)
    para(tf, v, size=12, color=MUTED, space_after=12)

rich(tf, [("One departure from the paper  ", True, RUST, False),
          ("I could not call a chondrocyte cluster: Acan is rarely detected "
           "(1.2% of all cells; highest cluster 7.05%). The Sox9-low E12.5 clusters "
           "look like early mesenchyme, so I labelled them that way rather than "
           "adopting their chondrocyte call.", False, INK, False)], size=12, space_after=0)
notes(s, "This is the baseline the chromatin claim must beat, and the annotation is "
         "independent of theirs \u2014 which is why the chondrocyte discrepancy is worth flagging.")

# ---------------------------------------------------------------- 5. story
s = slide()
header(s, "Step 3 \u00b7 the pilot", "RNA alone already predicts fate in \u201cunpatterned\u201d progenitors")
pic_fit(s, "figures/fig3_story.png", M, Inches(1.66), Inches(12.1), Inches(3.05))

y = Inches(5.05)
for i, (t, b) in enumerate([
    ("Trained on E14.5", "A logistic classifier separates anterior from posterior "
                         "cells at E14.5 almost perfectly (AUROC 0.999, n=4,259)."),
    ("Applied to E12.5", "On the 10,036 E12.5 cells with no A/P marker signature, "
                         "82% get a confident call (p<0.2 or p>0.8)."),
    ("Against a null", "With shuffled training labels, 0% are confident. The signal "
                       "is real, not a base-rate artifact.")]):
    xx = M + i * Inches(4.12)
    card(s, xx, y, Inches(3.88), Inches(1.30))
    tf = tb(s, xx + Inches(0.22), y + Inches(0.17), Inches(3.45), Inches(1.0))
    para(tf, t, size=14, bold=True, space_after=5, first=True)
    para(tf, b, size=12, color=MUTED, space_after=0)

c = card(s, M, Inches(6.52), CW, Inches(0.68), RGBColor(0xFC, 0xF3, 0xEC))
tf = tb(s, M + Inches(0.26), Inches(6.68), CW - Inches(0.5), Inches(0.5))
rich(tf, [("This complicates the hypothesis.  ", True, RUST, False),
          ("\u201cUnpatterned\u201d meant no marker signature, not no information. "
           "Chromatin priming now has to be shown to add predictive power over RNA \u2014 "
           "not merely to exist. That is a harder and more honest test.", False, INK, False)],
     size=13, space_after=0, first=True)
notes(s, "This was the surprise. I stress-tested it: removed the stage axis, balanced "
         "the classes, permuted the labels. The signal survives all three.")

# ---------------------------------------------------------------- 6. caveats
s = slide()
header(s, "How far to trust this", "What is verified, and what is not")

col = [("VERIFIED", GREEN, [
    "Both modalities present for all 9 libraries; sizes and pipeline version checked directly.",
    "QC reproduces the published thresholds to within ~1,900 cells, and the residual is explained.",
    "Peak-ID fragmentation measured exactly: 895,069 of 895,797 IDs appear in one library only.",
    "Marker/label concordance checked per cluster on absolute expression, not relative z-scores.",
    "Pilot classifier survives stage-axis removal, class balancing and a label-permutation null.",
]), ("NOT YET DONE", RUST, [
    "No ATAC analysis at all \u2014 the 12.3 GB fragment download and consensus matrix are pending.",
    "TSS-enrichment and nucleosome-signal QC not applied; they need the fragment files.",
    "Cluster labels are marker-based, not transferred from the authors' annotation.",
    "No batch integration beyond per-library HVG selection; residual library effects untested.",
    "The pilot uses RNA PCs, so it bounds the RNA baseline \u2014 it says nothing yet about chromatin.",
])]
for i, (title, color, items) in enumerate(col):
    xx = M + i * Inches(6.16)
    tf = tb(s, xx, Inches(1.72), Inches(5.8), Inches(4.6))
    para(tf, title, size=12, bold=True, color=color, space_after=13, first=True)
    for it in items:
        rich(tf, [("\u2014   ", False, color, False), (it, False, INK, False)],
             size=13, space_after=11)
notes(s, "The honest version: the feasibility question is settled, the biology question "
         "has not been touched yet.")

# ---------------------------------------------------------------- 7. next
s = slide()
header(s, "Next", "Three steps, in order")

rows = [("1", "Build the consensus ATAC matrix",
         "Download 12.3 GB of fragments, count all 38k cells against the 196,662 merged "
         "intervals, filter to peaks in \u22651% of cells, then TF-IDF/LSI dropping component 1.",
         "~3.2 GB in memory before filtering \u2014 needs care on a 7 GB machine."),
        ("2", "Test chromatin against the RNA baseline",
         "Same prediction task as the pilot, but with accessibility features: does chromatin "
         "at E12.5 predict E14.5 fate better than RNA does, at matched feature count and sparsity?",
         "The comparison is the result. A null result here is publishable as a negative."),
        ("3", "Robustness, then write-up",
         "Seeds, 50% subsampling, leave-one-timepoint-out; peak\u2013gene linkage and motif "
         "activity to nominate drivers; repo with config, tests and a Snakemake workflow.",
         "README must credit Yan et al. 2024 and Hussein et al. 2025 and state what neither did.")]
y = Inches(1.72)
for num, title, body, note in rows:
    card(s, M, y, CW, Inches(1.52))
    d = s.shapes.add_shape(9, M + Inches(0.30), y + Inches(0.40), Inches(0.52), Inches(0.52))
    d.fill.solid()
    d.fill.fore_color.rgb = GREEN
    d.line.fill.background()
    d.shadow.inherit = False
    dtf = d.text_frame
    dtf.text = num
    dtf.paragraphs[0].alignment = PP_ALIGN.CENTER
    dtf.paragraphs[0].runs[0].font.size = Pt(18)
    dtf.paragraphs[0].runs[0].font.bold = True
    dtf.paragraphs[0].runs[0].font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
    tf = tb(s, M + Inches(1.05), y + Inches(0.20), Inches(10.6), Inches(1.15))
    para(tf, title, size=15, bold=True, space_after=4, first=True)
    para(tf, body, size=12.5, color=MUTED, space_after=4)
    rich(tf, [("Watch out:  ", True, RUST, False), (note, False, MUTED, False)],
         size=11.5, space_after=0)
    y += Inches(1.66)

tf = tb(s, M, Inches(6.86), CW, Inches(0.35))
para(tf, "Timeline: the original 10-day estimate is optimistic by roughly half, "
         "mostly in step 1 and in model training.", size=12, color=MUTED, space_after=0, first=True)
notes(s, "Step 2 is the project. Steps 1 and 3 are the cost of doing it properly.")

prs.save("docs/status_deck.pptx")
print(f"wrote docs/status_deck.pptx ({len(prs.slides)} slides)")
